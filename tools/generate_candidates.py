# ---------------------------------------------------------------------------
# generate_candidates.py — מחולל את קובץ "מועמדים בהליך" (טבלה 2 באפיון).
#
# מפיק שני פלטים משורש אמת אחד:
#   1. הקובץ למסירה — מועמדים_בהליך_<תאריך>.xlsx בשורש הפרויקט, בדיוק בפורמט
#      שהמערכת החדשה אמורה לקבל: מחוז · תחנה · כמות מועמדים.
#   2. approved_candidates.json (בתיקיית הכלים) — שממנו ההדגמה טוענת את
#      המספרים, כך שהמפה/הטבלה/הייצוא מציגים בדיוק את אותם נתונים.
#
# חלוקה: סך 2000 מועמדים על פני 88 התחנות, בשיטת "השארית הגדולה" (סכום
# מדויק). המשקל לכל תחנה = תקנים חסרים + רצפה — כדי שהמועמדים יתרכזו היכן
# שהפער גדול (שם מתנהל הגיוס), אך לכל תחנה יש בסיס (תחלופה שוטפת).
#
# הרצה:  python to_github/tools/generate_candidates.py
# ---------------------------------------------------------------------------
import json
import math
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
APPROVED = HERE / "approved_stations.json"
OUT_JSON = HERE / "approved_candidates.json"
OUT_XLSX = HERE.parent.parent / f"מועמדים_בהליך_{date.today().isoformat()}.xlsx"

TOTAL = 2000
FLOOR = 2  # רצפת מועמדים לכל תחנה (גיוס שוטף גם בתחנה מלאה)


def _distribute(rows):
    """מחלק TOTAL בין השורות לפי משקל, כך שהסכום מדויק == TOTAL."""
    total_w = sum(r["weight"] for r in rows) or 1
    raw = [TOTAL * r["weight"] / total_w for r in rows]
    counts = [int(math.floor(x)) for x in raw]
    remainder = TOTAL - sum(counts)
    order = sorted(range(len(rows)), key=lambda i: raw[i] - counts[i], reverse=True)
    for i in order[:remainder]:
        counts[i] += 1
    for r, c in zip(rows, counts):
        r["count"] = c


def main():
    data = json.loads(APPROVED.read_text(encoding="utf-8"))
    rows = []
    for name, s in data.items():
        req, act = s.get("required"), s.get("actual")
        missing = max(0, req - act) if (req is not None and act is not None) else 0
        rows.append({"district": s["district"], "name": name, "weight": missing + FLOOR})

    _distribute(rows)
    assert sum(r["count"] for r in rows) == TOTAL, "הסכום חייב להיות בדיוק 2000"

    # (1) JSON לצריכת ההדגמה
    OUT_JSON.write_text(
        json.dumps({r["name"]: r["count"] for r in rows}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )

    # (2) קובץ אקסל למסירה — מחוז · תחנה · כמות מועמדים
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "מועמדים בהליך"
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin", color="C8CDD6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="14284B")
    head_font = Font(name="David", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="David", size=11)
    center = Alignment(horizontal="center", vertical="center")

    for c, h in enumerate(["מחוז", "תחנה", "כמות מועמדים"], 1):
        cell = ws.cell(1, c, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = border

    r = 2
    for row in sorted(rows, key=lambda x: (x["district"], x["name"])):
        for c, v in enumerate([row["district"], row["name"], row["count"]], 1):
            cell = ws.cell(r, c, v)
            cell.font = cell_font
            cell.alignment = center
            cell.border = border
        r += 1

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 15
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)

    print(f"נכתבו {len(rows)} תחנות · סה\"כ {TOTAL} מועמדים.")
    print(f"  קובץ למסירה: {OUT_XLSX.name}")
    print(f"  להדגמה:      {OUT_JSON.name}")


if __name__ == "__main__":
    main()

# ---------------------------------------------------------------------------
# build_export_xlsx.py — בונה את קובץ הייצוא לאקסל של ההדגמה.
#
# מה זה
# -----
# בגרסת השרת המלאה כפתור "ייצוא אקסל" הפיק קובץ מהנתונים החיים. בגרסת
# הפורטפוליו הסטטית (GitHub Pages) אין שרת, ולכן הקובץ מיוצר מראש מתוך
# data/*.json (אותם קבצים שה-JS טוען) ונשמר כ-data/export.xlsx. כפתורי
# הייצוא באתר מורידים בדיוק את הקובץ הזה.
#
# פרטיות
# ------
# הקובץ נבנה אך ורק מ-data/*.json — נתונים סינתטיים שהוגרלו ב-seed קבוע.
# אין בו, ולא נגזר ממנו, שום נתון תפעולי אמיתי. זה הקובץ היחיד מסוג *.xlsx
# שמותר במאגר הציבורי (ראו החריג ב-.gitignore).
#
# מבנה הטבלה (משקף את קובץ המקור, מסודר לארבע משפחות מקצוע)
# --------------------------------------------------------
# כותרת דו-שורתית: קבוצה ← שדות. לכל תחנה, ולכל אחת מ-4 המשפחות
# (חוקר/סייר/בלש/אחר): תקנים · תקנים פנויים · אחוז איוש. בנוסף סה"כ-תחנה
# ודמוגרפיה (% גברים, % יהודים). גיליון שני = סיכום (כמויות + התפלגות).
# ---------------------------------------------------------------------------
import json
import statistics
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FAMILIES = ["חוקר", "סייר", "בלש", "אחר"]

_THIN = Side(style="thin", color="C8CDD6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_GRP_FILL = PatternFill("solid", fgColor="111827")
_SUB_FILL = PatternFill("solid", fgColor="374151")
_WHITE_B = Font(name="David", bold=True, color="FFFFFF", size=11)
_CELL = Font(name="David", size=11)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _status_color(pct):
    if pct is None:
        return "9CA3AF"
    if pct < 70:
        return "DC2626"
    if pct < 80:
        return "EA580C"
    if pct < 90:
        return "EAB308"
    return "16A34A"


def _load(data_dir):
    data_dir = Path(data_dir)
    stations = json.loads((data_dir / "stations.json").read_text(encoding="utf-8"))
    details = {
        s["id"]: json.loads((data_dir / "station-detail" / f"{s['id']}.json").read_text(encoding="utf-8"))
        for s in stations
    }
    return stations, details


def _sheet_stations(wb, stations, details):
    ws = wb.active
    ws.title = "תחנות ומקצועות"
    ws.sheet_view.rightToLeft = True

    groups = [("זיהוי", ["מחוז", "מרחב", "תחנה"]),
              ('סה"כ תחנה', ["תקנים", "תקנים פנויים", "אחוז איוש"])]
    for fam in FAMILIES:
        groups.append((fam, ["תקנים", "תקנים פנויים", "אחוז איוש"]))
    groups.append(("דמוגרפיה (תחנה)", ["% גברים", "% יהודים"]))
    groups.append(("מועמדים", ["בהליך"]))

    col = 1
    for gname, subs in groups:
        start = col
        for sub in subs:
            ws.cell(2, col, sub)
            col += 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
        ws.cell(1, start, gname)
    ncols = col - 1
    for c in range(1, ncols + 1):
        g = ws.cell(1, c); g.fill = _GRP_FILL; g.font = _WHITE_B; g.alignment = _CENTER; g.border = _BORDER
        s = ws.cell(2, c); s.fill = _SUB_FILL; s.font = _WHITE_B; s.alignment = _CENTER; s.border = _BORDER

    rows = sorted(stations, key=lambda s: (s["district"], s["area"], s["name"]))
    r = 3
    for s in rows:
        by_label = {x["label"]: x for x in details[s["id"]]["roles"]}
        vals = [s["district"], s["area"], s["name"],
                s["required_positions"], s["missing_positions"], s["staffing_pct"] / 100.0]
        for fam in FAMILIES:
            x = by_label[fam]
            vals += [x["required_positions"], x["missing"], (x["staffing_pct"] or 0) / 100.0]
        vals += [s["pct_male"] / 100.0, s["pct_jewish"] / 100.0]
        # מועמדים בהליך — כלל הברזל: None = "אין נתון", לא 0.
        cand = s.get("candidates_in_process")
        vals.append(cand if cand is not None else "אין נתון")
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.font = _CELL; cell.alignment = _CENTER; cell.border = _BORDER
        # עמודות אחוז-איוש: פורמט אחוז + צביעה לפי סטטוס (סה"כ ולכל משפחה)
        pct_cols = {6: s["staffing_pct"]}
        ci = 6
        for fam in FAMILIES:
            ci += 3
            pct_cols[ci] = by_label[fam]["staffing_pct"]
        for c, p in pct_cols.items():
            cell = ws.cell(r, c); cell.number_format = "0.0%"
            cell.fill = PatternFill("solid", fgColor=_status_color(p))
            cell.font = Font(name="David", color="FFFFFF", size=11, bold=(c == 6))
        # עמודות הדמוגרפיה (% גברים/יהודים) — מיקום קבוע, לפני עמודת המועמדים.
        demo_first = 3 + 3 + 3 * len(FAMILIES) + 1
        ws.cell(r, demo_first).number_format = "0%"
        ws.cell(r, demo_first + 1).number_format = "0%"
        r += 1

    widths = [11, 11, 20] + [9, 11, 10] * (1 + len(FAMILIES)) + [10, 10, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D3"
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 30


def _sheet_summary(wb, stations, details):
    ws = wb.create_sheet("סיכום")
    ws.sheet_view.rightToLeft = True

    def H(txt, row):
        ws.cell(row, 1, txt).font = Font(name="David", bold=True, size=13)
        return row + 1

    def HR(row, cols):
        for j, t in enumerate(cols, 1):
            c = ws.cell(row, j, t); c.fill = _GRP_FILL; c.font = _WHITE_B; c.alignment = _CENTER; c.border = _BORDER
        return row + 1

    def DR(row, vals, fill=None):
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v); c.font = _CELL; c.alignment = _CENTER; c.border = _BORDER
            if fill and j == 1:
                c.fill = PatternFill("solid", fgColor=fill)
                c.font = Font(name="David", bold=True, color="FFFFFF", size=11)
        return row + 1

    pcts = sorted(s["staffing_pct"] for s in stations)
    r = 1
    r = H("כמויות", r)
    r = HR(r, ["מדד", "כמות"])
    r = DR(r, ["מחוזות", len({s["district"] for s in stations})])
    r = DR(r, ["מרחבים", len({(s["district"], s["area"]) for s in stations})])
    r = DR(r, ["תחנות", len(stations)])
    r += 1

    r = H("התפלגות אחוז איוש (עקומת פעמון)", r)
    ws.cell(r, 1, f"ממוצע {statistics.mean(pcts):.1f}%  ·  חציון {statistics.median(pcts):.0f}%  ·  "
                  f"סטיית תקן {statistics.pstdev(pcts):.1f}").font = _CELL
    r += 1
    r = HR(r, ["טווח", "תחנות", "גרף"])
    buckets = Counter(int(p // 5) * 5 for p in pcts)
    for k in range(50, 101, 5):
        n = buckets.get(k, 0)
        r = DR(r, [f"{k}-{k + 4}%", n, "█" * n])
    r += 1

    r = H("פילוח לפי סטטוס", r)
    r = HR(r, ["סטטוס", "תחנות"])
    sc = Counter(s["status"] for s in stations)
    ref = {"קריטי": 60, "דחוף": 75, "בינוני": 85, "תקין": 95}
    for stt in ["קריטי", "דחוף", "בינוני", "תקין"]:
        r = DR(r, [stt, sc.get(stt, 0)], fill=_status_color(ref[stt]))
    r += 1

    r = H("פילוח לפי מחוז", r)
    r = HR(r, ["מחוז", "תחנות", "מרחבים", 'סה"כ תקן', 'סה"כ פנוי', "אחוז איוש"])
    byd = defaultdict(lambda: {"n": 0, "areas": set(), "req": 0, "act": 0, "vac": 0})
    for s in stations:
        d = byd[s["district"]]
        d["n"] += 1; d["areas"].add(s["area"])
        d["req"] += s["required_positions"]; d["act"] += s["actual_positions"]; d["vac"] += s["missing_positions"]
    for dist, d in sorted(byd.items(), key=lambda x: -x[1]["n"]):
        r2 = DR(r, [dist, d["n"], len(d["areas"]), d["req"], d["vac"], round(d["act"] / d["req"], 3)])
        ws.cell(r, 6).number_format = "0.0%"; r = r2
    r += 1

    r = H("פילוח לפי משפחת מקצוע", r)
    r = HR(r, ["משפחה", 'סה"כ תקן', 'סה"כ פנוי', "אחוז איוש"])
    fam = defaultdict(lambda: {"req": 0, "act": 0, "vac": 0})
    for d in details.values():
        for x in d["roles"]:
            f = fam[x["label"]]
            f["req"] += x["required_positions"]; f["act"] += x["actual_positions"]; f["vac"] += x["missing"]
    for name in FAMILIES:
        f = fam[name]
        r2 = DR(r, [name, f["req"], f["vac"], round(f["act"] / f["req"], 3)])
        ws.cell(r, 4).number_format = "0.0%"; r = r2

    for i, w in enumerate([16, 10, 12, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(data_dir, out_path=None):
    """בונה את data/export.xlsx מתוך קובצי ה-JSON הסינתטיים. מחזיר את הנתיב."""
    data_dir = Path(data_dir)
    out_path = Path(out_path) if out_path else data_dir / "export.xlsx"
    stations, details = _load(data_dir)
    wb = Workbook()
    _sheet_stations(wb, stations, details)
    _sheet_summary(wb, stations, details)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    p = build(Path(__file__).resolve().parent.parent / "data")
    print(f"נכתב {p}")
